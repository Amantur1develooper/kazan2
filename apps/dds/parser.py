"""
DDS (ДДС) Excel parser for 1С cash flow export.

Expected columns (row 0 = header):
  0: Дата
  1: Банк / Касса
  2: Операция
  3: Сумма            (positive = income, negative = expense, empty for transfers)
  4: Сумма перемещения (for transfers)
  5: Откуда / куда
  6: Касса / счет
  7: Объект
  8: Назначение
"""

import re
from datetime import date
from decimal import Decimal, InvalidOperation


def _to_decimal(value) -> Decimal:
    if value is None or value == '':
        return Decimal('0')
    s = str(value).strip().replace('\xa0', '').replace(' ', '').replace(',', '.')
    s = re.sub(r'[^\d.\-]', '', s)
    if not s or s in ('-', '.'):
        return Decimal('0')
    try:
        return Decimal(s)
    except InvalidOperation:
        return Decimal('0')


def _parse_date_cell(ws, r, c, datemode):
    import xlrd
    cell = ws.cell(r, c)
    if cell.ctype == xlrd.XL_CELL_DATE:
        t = xlrd.xldate_as_tuple(cell.value, datemode)
        try:
            return date(t[0], t[1], t[2])
        except ValueError:
            return None
    s = str(cell.value).strip()
    m = re.search(r'(\d{1,2})\.(\d{1,2})\.(\d{4})', s)
    if m:
        try:
            return date(int(m.group(3)), int(m.group(2)), int(m.group(1)))
        except ValueError:
            return None
    return None


def _parse_date_value(val) -> 'date | None':
    """Parse a date from an openpyxl cell value (datetime, date, or string)."""
    if val is None:
        return None
    from datetime import datetime as _dt, date as _date
    if isinstance(val, _dt):
        return val.date()
    if isinstance(val, _date):
        return val
    s = str(val).strip()
    m = re.search(r'(\d{1,2})\.(\d{1,2})\.(\d{4})', s)
    if m:
        try:
            return date(int(m.group(3)), int(m.group(2)), int(m.group(1)))
        except ValueError:
            return None
    return None


def _str(v) -> str:
    if v is None:
        return ''
    return str(v).strip()


def _has_transfer_col(header_row) -> bool:
    """Return True if the header contains a 'Сумма перемещения' column."""
    return any('перемещен' in str(h).lower() for h in header_row if h)


def _extract_row_fields(values, has_transfer):
    """
    Extract (amount_raw, transfer_raw, counterparty, account_name, object_ref, description)
    from a value tuple, adjusting for whether the transfer column is present.
    """
    n = len(values)
    amount_raw = _to_decimal(values[3] if n > 3 else None)
    if has_transfer:
        transfer_raw  = _to_decimal(values[4] if n > 4 else None)
        counterparty  = _str(values[5] if n > 5 else None)
        account_name  = _str(values[6] if n > 6 else None)
        object_ref    = _str(values[7] if n > 7 else None)
        description   = _str(values[8] if n > 8 else None)
    else:
        transfer_raw  = Decimal('0')
        counterparty  = _str(values[4] if n > 4 else None)
        account_name  = _str(values[5] if n > 5 else None)
        object_ref    = _str(values[6] if n > 6 else None)
        description   = _str(values[7] if n > 7 else None)
    return amount_raw, transfer_raw, counterparty, account_name, object_ref, description


def _parse_dds_xls(filepath: str) -> list:
    import xlrd
    wb = xlrd.open_workbook(filepath)
    ws = wb.sheet_by_index(0)
    header = [ws.cell_value(0, c) for c in range(ws.ncols)]
    has_transfer = _has_transfer_col(header)
    records = []
    for r in range(1, ws.nrows):
        row_date = _parse_date_cell(ws, r, 0, wb.datemode)
        if not row_date:
            continue
        bank_or_cash = _str(ws.cell_value(r, 1))
        op_type = _str(ws.cell_value(r, 2))
        if not op_type:
            continue
        values = tuple(ws.cell_value(r, c) for c in range(ws.ncols))
        amount_raw, transfer_raw, counterparty, account_name, object_ref, description = \
            _extract_row_fields(values, has_transfer)
        _append_record(records, row_date, bank_or_cash, op_type, amount_raw, transfer_raw,
                       counterparty, account_name, object_ref, description)
    return records


def _parse_dds_xlsx(filepath: str) -> list:
    import openpyxl
    wb = openpyxl.load_workbook(filepath, read_only=True, data_only=True)
    ws = wb.worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    if not rows:
        return []
    has_transfer = _has_transfer_col(rows[0])
    records = []
    for row in rows[1:]:
        row_date = _parse_date_value(row[0] if row else None)
        if not row_date:
            continue
        bank_or_cash = _str(row[1] if len(row) > 1 else None)
        op_type = _str(row[2] if len(row) > 2 else None)
        if not op_type:
            continue
        amount_raw, transfer_raw, counterparty, account_name, object_ref, description = \
            _extract_row_fields(row, has_transfer)
        _append_record(records, row_date, bank_or_cash, op_type, amount_raw, transfer_raw,
                       counterparty, account_name, object_ref, description)
    return records


def _append_record(records, row_date, bank_or_cash, op_type, amount_raw, transfer_raw,
                   counterparty, account_name, object_ref, description):
    op_lower = op_type.lower()
    if 'перемещение' in op_lower or (transfer_raw != 0 and amount_raw == 0):
        direction = 'transfer'
        amount = abs(transfer_raw)
    elif amount_raw > 0:
        direction = 'in'
        amount = amount_raw
    elif amount_raw < 0:
        direction = 'out'
        amount = abs(amount_raw)
    else:
        return
    records.append({
        'operation_date': row_date,
        'bank_or_cash': bank_or_cash,
        'operation_type': op_type,
        'direction': direction,
        'amount': amount,
        'transfer_amount': abs(transfer_raw),
        'counterparty': counterparty,
        'account_name': account_name,
        'object_ref': object_ref,
        'description': description,
    })


def parse_dds_excel(filepath: str) -> list:
    """Parse 1С ДДС Excel file (.xls or .xlsx). Returns list of record dicts."""
    if filepath.lower().endswith('.xlsx'):
        return _parse_dds_xlsx(filepath)
    return _parse_dds_xls(filepath)


def auto_match_block(object_ref: str, blocks) -> object:
    """Try to find a Block whose name appears in object_ref string."""
    if not object_ref:
        return None
    obj_lower = object_ref.lower()
    for block in sorted(blocks, key=lambda b: len(b.name), reverse=True):
        if block.name.lower() in obj_lower:
            return block
    return None


def apply_dds_import(residential_complex, records: list, dds_import, default_block=None) -> dict:
    """
    Persist parsed DDS records to the database.
    Auto-creates CashAccount records per unique account_name.
    If default_block is set, all records get that block; otherwise auto-matches by object_ref.
    Returns stats dict.
    """
    from .models import CashAccount, CashFlowRecord
    from apps.projects.models import Block

    # Guard: skip if already applied
    if dds_import and CashFlowRecord.objects.filter(source_import=dds_import).exists():
        return {'created': 0, 'errors': 0}

    # Pre-load accounts cache: name → CashAccount
    account_cache = {}
    for acc in CashAccount.objects.filter(residential_complex=residential_complex):
        account_cache[acc.name] = acc

    # Pre-load all blocks for auto-matching (only needed when no default_block)
    all_blocks = [] if default_block is not None else list(
        Block.objects.filter(residential_complex=residential_complex).select_related('residential_complex')
    )

    stats = {'created': 0, 'errors': 0}
    to_create = []

    for item in records:
        try:
            acc_name = item.get('account_name', '')
            if acc_name and acc_name not in account_cache:
                acc_type = 'bank' if 'банк' in item.get('bank_or_cash', '').lower() or any(
                    c.isdigit() for c in acc_name[:6]
                ) else 'cash'
                acc = CashAccount.objects.create(
                    name=acc_name,
                    account_type=acc_type,
                    residential_complex=residential_complex,
                    block=default_block,
                )
                account_cache[acc_name] = acc
            account = account_cache.get(acc_name)

            block = default_block if default_block is not None else auto_match_block(item.get('object_ref', ''), all_blocks)

            op_date = item['operation_date']
            if hasattr(op_date, 'isoformat'):
                pass  # already a date object
            else:
                from datetime import datetime
                op_date = datetime.strptime(op_date, '%Y-%m-%d').date()

            to_create.append(CashFlowRecord(
                operation_date=op_date,
                bank_or_cash=item.get('bank_or_cash', ''),
                operation_type=item.get('operation_type', ''),
                direction=item.get('direction', 'out'),
                amount=item.get('amount', 0),
                transfer_amount=item.get('transfer_amount', 0),
                counterparty=item.get('counterparty', ''),
                account_name=acc_name,
                object_ref=item.get('object_ref', ''),
                description=item.get('description', ''),
                account=account,
                block=block,
                source_import=dds_import,
            ))
            stats['created'] += 1
        except Exception:
            stats['errors'] += 1

    CashFlowRecord.objects.bulk_create(to_create)
    return stats
